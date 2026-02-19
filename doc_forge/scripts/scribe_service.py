#!/usr/bin/env python3
"""Eidosian Doc Processor v2.

Production-oriented documentation processor with:
- Multi-format extraction (code/config/text/pdf/docx/html/xml/csv/...)
- Staging -> federated judge gate -> finalization
- Resumable persistent state with atomic writes
- Realtime status API + dashboard UI
- Optional managed llama.cpp server startup
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import threading
import time
from collections import Counter, deque
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
from fastapi import FastAPI, Query
from fastapi.responses import HTMLResponse, JSONResponse
from pydantic import BaseModel
import uvicorn

try:
    from bs4 import BeautifulSoup
except Exception:  # pragma: no cover
    BeautifulSoup = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None


SUPPORTED_SUFFIXES = {
    ".py", ".pyi", ".js", ".jsx", ".ts", ".tsx", ".java", ".kt", ".kts", ".go", ".rs", ".c", ".h", ".hpp", ".cc", ".cpp",
    ".cs", ".swift", ".rb", ".php", ".lua", ".sh", ".bash", ".zsh", ".ps1", ".sql",
    ".md", ".rst", ".txt", ".adoc", ".org", ".log",
    ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".xml", ".csv", ".tsv", ".env",
    ".html", ".htm", ".xhtml", ".css", ".scss", ".less", ".svg",
    ".pdf", ".docx",
}

PLACEHOLDER_MARKERS = {
    "todo",
    "lorem ipsum",
    "placeholder",
    "i cannot",
    "i'm unable",
    "insufficient context",
    "not provided",
    "not available",
    "unknown",
}

COMMON_STOPWORDS = {
    "the", "and", "for", "that", "this", "with", "from", "into", "your", "file", "code", "data", "are", "was", "were", "has", "have", "had",
    "will", "would", "should", "could", "can", "not", "but", "you", "our", "their", "its", "then", "than", "when", "where", "what", "which",
    "there", "here", "about", "over", "under", "also", "while", "using", "used", "use", "each", "per", "all", "any", "may", "one", "two",
    "three", "four", "five", "via", "out", "in", "on", "as", "at", "to", "of", "is", "it", "by", "an", "or", "be", "if", "do", "does",
}

REQUIRED_HEADINGS = [
    "# File Overview",
    "## Key Structures",
    "## Behavior Summary",
    "## Validation Notes",
    "## Improvement Opportunities",
]


class StatusPayload(BaseModel):
    status: str
    started_at: str
    model: str
    total_discovered: int
    processed: int
    approved: int
    rejected: int
    errors: int
    remaining: int
    average_quality_score: float
    average_seconds_per_document: float
    eta_seconds: float | None
    last_staged: str | None
    last_approved: str | None
    top_tags: list[tuple[str, int]]
    top_themes: list[tuple[str, int]]
    doc_type_frequency: list[tuple[str, int]]


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    tmp.replace(path)


def _read_json(path: Path, default: dict[str, Any]) -> dict[str, Any]:
    if not path.exists():
        return default
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default
    if not isinstance(data, dict):
        return default
    return data


def _registry_port(forge_root: Path, service_key: str, fallback: int) -> int:
    registry_path = forge_root / "config" / "ports.json"
    if not registry_path.exists():
        return fallback
    try:
        payload = json.loads(registry_path.read_text(encoding="utf-8"))
    except Exception:
        return fallback
    services = payload.get("services") if isinstance(payload, dict) else None
    if not isinstance(services, dict):
        return fallback
    service = services.get(service_key)
    if not isinstance(service, dict):
        return fallback
    try:
        port = int(service.get("port", fallback))
    except Exception:
        return fallback
    return port if port > 0 else fallback


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _is_text_bytes(sample: bytes) -> bool:
    if not sample:
        return True
    if b"\x00" in sample:
        return False
    text_bytes = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(0x20, 0x100)))
    return all(ch in text_bytes for ch in sample)


def _extract_terms(text: str, limit: int = 24) -> list[str]:
    words = re.findall(r"[A-Za-z][A-Za-z0-9_\-]{3,}", text.lower())
    words = [w for w in words if w not in COMMON_STOPWORDS]
    freq = Counter(words)
    return [w for w, _ in freq.most_common(limit)]


def _extract_symbols(source_text: str) -> list[str]:
    patterns = [
        r"\bdef\s+([A-Za-z_][A-Za-z0-9_]*)",
        r"\bclass\s+([A-Za-z_][A-Za-z0-9_]*)",
        r"\bfunction\s+([A-Za-z_][A-Za-z0-9_]*)",
        r"\bconst\s+([A-Za-z_][A-Za-z0-9_]*)",
        r"\blet\s+([A-Za-z_][A-Za-z0-9_]*)",
        r"\binterface\s+([A-Za-z_][A-Za-z0-9_]*)",
        r"\benum\s+([A-Za-z_][A-Za-z0-9_]*)",
    ]
    out: list[str] = []
    for pattern in patterns:
        out.extend(re.findall(pattern, source_text))
    unique: list[str] = []
    seen: set[str] = set()
    for item in out:
        if item not in seen:
            unique.append(item)
            seen.add(item)
    return unique[:80]


@dataclass
class ProcessorConfig:
    forge_root: Path
    source_root: Path
    runtime_root: Path
    staging_root: Path
    final_root: Path
    rejected_root: Path
    judgments_root: Path
    index_path: Path
    state_path: Path
    status_path: Path
    host: str
    port: int
    completion_url: str
    llm_model_path: Path
    llama_server_bin: Path
    llm_server_port: int
    llm_ctx_size: int
    llm_temperature: float
    llm_n_predict: int
    include_suffixes: set[str]
    max_file_bytes: int
    max_chars: int
    loop_sleep_s: float
    approval_threshold: float
    min_grounding_overlap: float
    enable_managed_llm: bool
    dry_run: bool

    @classmethod
    def from_env(cls, forge_root: Path, host: str, port: int, dry_run: bool = False) -> "ProcessorConfig":
        runtime_root = forge_root / "doc_forge" / "runtime"
        staging_root = runtime_root / "staging_docs"
        final_root = runtime_root / "final_docs"
        rejected_root = runtime_root / "rejected_docs"
        judgments_root = runtime_root / "judgments"

        model_path = Path(
            os.environ.get(
                "EIDOS_DOC_FORGE_MODEL",
                "models/Qwen2.5-0.5B-Instruct-Q8_0.gguf",
            )
        )
        if not model_path.is_absolute():
            model_path = forge_root / model_path

        default_llm_port = _registry_port(forge_root, "doc_forge_llm", 8093)
        completion_url = os.environ.get(
            "EIDOS_DOC_FORGE_COMPLETION_URL",
            f"http://127.0.0.1:{os.environ.get('EIDOS_DOC_FORGE_LLM_PORT', str(default_llm_port))}/completion",
        )

        llama_server_bin = Path(
            os.environ.get("EIDOS_DOC_FORGE_LLAMA_SERVER_BIN", "llama.cpp/build/bin/llama-server")
        )
        if not llama_server_bin.is_absolute():
            llama_server_bin = forge_root / llama_server_bin

        source_root = Path(os.environ.get("EIDOS_DOC_FORGE_SOURCE_ROOT", str(forge_root)))
        if not source_root.is_absolute():
            source_root = forge_root / source_root

        return cls(
            forge_root=forge_root,
            source_root=source_root,
            runtime_root=runtime_root,
            staging_root=staging_root,
            final_root=final_root,
            rejected_root=rejected_root,
            judgments_root=judgments_root,
            index_path=runtime_root / "doc_index.json",
            state_path=runtime_root / "processor_state.json",
            status_path=runtime_root / "processor_status.json",
            host=host,
            port=port,
            completion_url=completion_url,
            llm_model_path=model_path,
            llama_server_bin=llama_server_bin,
            llm_server_port=int(os.environ.get("EIDOS_DOC_FORGE_LLM_PORT", str(default_llm_port))),
            llm_ctx_size=int(os.environ.get("EIDOS_DOC_FORGE_CTX_SIZE", "4096")),
            llm_temperature=float(os.environ.get("EIDOS_DOC_FORGE_TEMPERATURE", "0.2")),
            llm_n_predict=int(os.environ.get("EIDOS_DOC_FORGE_N_PREDICT", "2200")),
            include_suffixes=set(SUPPORTED_SUFFIXES),
            max_file_bytes=int(os.environ.get("EIDOS_DOC_FORGE_MAX_FILE_BYTES", str(4 * 1024 * 1024))),
            max_chars=int(os.environ.get("EIDOS_DOC_FORGE_MAX_CHARS", "45000")),
            loop_sleep_s=float(os.environ.get("EIDOS_DOC_FORGE_LOOP_SLEEP_S", "0.2")),
            approval_threshold=float(os.environ.get("EIDOS_DOC_FORGE_APPROVAL_THRESHOLD", "0.78")),
            min_grounding_overlap=float(os.environ.get("EIDOS_DOC_FORGE_MIN_GROUNDING_OVERLAP", "0.08")),
            enable_managed_llm=os.environ.get("EIDOS_DOC_FORGE_ENABLE_MANAGED_LLM", "1") == "1",
            dry_run=dry_run,
        )


class ManagedModelServer:
    def __init__(self, cfg: ProcessorConfig) -> None:
        self.cfg = cfg
        self.proc: subprocess.Popen[Any] | None = None
        self.log_handle: Any = None

    def _http_ready(self, timeout_s: float = 1.5) -> bool:
        try:
            r = requests.get(self.cfg.completion_url.replace("/completion", "/health"), timeout=timeout_s)
            return r.status_code < 500
        except Exception:
            return False

    def is_ready(self, timeout_s: float = 1.0) -> bool:
        return self._http_ready(timeout_s=timeout_s)

    def start(self) -> None:
        if self._http_ready():
            return
        if not self.cfg.enable_managed_llm:
            raise RuntimeError("LLM completion endpoint unavailable and managed startup disabled")
        if not self.cfg.llama_server_bin.exists():
            raise RuntimeError(f"Missing llama-server binary: {self.cfg.llama_server_bin}")
        if not self.cfg.llm_model_path.exists():
            raise RuntimeError(f"Missing doc processor model: {self.cfg.llm_model_path}")

        self.cfg.runtime_root.mkdir(parents=True, exist_ok=True)
        log_path = self.cfg.runtime_root / "llama_server.log"
        self.log_handle = log_path.open("a")
        cmd = [
            str(self.cfg.llama_server_bin),
            "-m",
            str(self.cfg.llm_model_path),
            "--port",
            str(self.cfg.llm_server_port),
            "--ctx-size",
            str(self.cfg.llm_ctx_size),
            "--temp",
            str(self.cfg.llm_temperature),
            "--n-gpu-layers",
            os.environ.get("EIDOS_LLAMA_GPU_LAYERS", "0"),
            "--parallel",
            os.environ.get("EIDOS_DOC_FORGE_PARALLEL", "1"),
        ]

        env = os.environ.copy()
        bin_dir = str(self.cfg.llama_server_bin.resolve().parent)
        env["PATH"] = f"{bin_dir}:{env.get('PATH', '')}"
        ld = env.get("LD_LIBRARY_PATH", "")
        env["LD_LIBRARY_PATH"] = f"{bin_dir}:{ld}" if ld else bin_dir

        self.proc = subprocess.Popen(cmd, stdout=self.log_handle, stderr=subprocess.STDOUT, env=env)
        start = time.time()
        while time.time() - start < 90:
            if self._http_ready(timeout_s=2.0):
                return
            if self.proc.poll() is not None:
                raise RuntimeError(f"llama-server exited with code {self.proc.returncode}")
            time.sleep(0.5)
        raise TimeoutError("Timed out waiting for managed llama-server")

    def stop(self) -> None:
        if self.proc and self.proc.poll() is None:
            self.proc.terminate()
            try:
                self.proc.wait(timeout=8)
            except subprocess.TimeoutExpired:
                self.proc.kill()
        self.proc = None
        if self.log_handle is not None:
            self.log_handle.close()
            self.log_handle = None


class DocumentExtractor:
    def __init__(self, cfg: ProcessorConfig) -> None:
        self.cfg = cfg

    def extract(self, path: Path) -> tuple[str, dict[str, Any]]:
        suffix = path.suffix.lower()
        if suffix in {".pdf"}:
            return self._extract_pdf(path), {"doc_type": "pdf"}
        if suffix in {".docx"}:
            return self._extract_docx(path), {"doc_type": "docx"}
        if suffix in {".html", ".htm", ".xhtml", ".svg", ".xml"}:
            return self._extract_html_like(path), {"doc_type": suffix.lstrip(".")}
        return self._extract_text(path), {"doc_type": suffix.lstrip(".") or "text"}

    def _extract_text(self, path: Path) -> str:
        raw = path.read_bytes()
        if len(raw) > self.cfg.max_file_bytes:
            raw = raw[: self.cfg.max_file_bytes]
        if not _is_text_bytes(raw[:4096]):
            raise ValueError("binary-like content")
        text = raw.decode("utf-8", errors="replace")
        return text[: self.cfg.max_chars]

    def _extract_pdf(self, path: Path) -> str:
        if PdfReader is None:
            raise RuntimeError("pypdf unavailable")
        reader = PdfReader(str(path))
        parts: list[str] = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
            if sum(len(p) for p in parts) > self.cfg.max_chars:
                break
        text = "\n\n".join(parts)
        return text[: self.cfg.max_chars]

    def _extract_docx(self, path: Path) -> str:
        if DocxDocument is None:
            raise RuntimeError("python-docx unavailable")
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        return text[: self.cfg.max_chars]

    def _extract_html_like(self, path: Path) -> str:
        text = path.read_text(encoding="utf-8", errors="replace")
        if BeautifulSoup is None:
            return text[: self.cfg.max_chars]
        soup = BeautifulSoup(text, "html.parser")
        body_text = soup.get_text("\n", strip=True)
        return body_text[: self.cfg.max_chars]


class DocGenerator:
    def __init__(self, cfg: ProcessorConfig) -> None:
        self.cfg = cfg

    def generate(self, rel_path: str, source_text: str, metadata: dict[str, Any]) -> str:
        base_prompt = self._build_prompt(rel_path=rel_path, source_text=source_text, metadata=metadata)
        prompt = base_prompt
        last_error: Exception | None = None
        for attempt in range(1, 4):
            payload = {
                "prompt": prompt,
                "n_predict": self.cfg.llm_n_predict,
                "temperature": self.cfg.llm_temperature,
                "stop": ["<|im_end|>", "</s>"],
                "stream": False,
            }
            resp = requests.post(self.cfg.completion_url, json=payload, timeout=240)
            resp.raise_for_status()
            content = (resp.json().get("content") or "").strip()
            if not content:
                last_error = RuntimeError("empty model output")
            else:
                try:
                    return self._normalize(content)
                except ValueError as exc:
                    last_error = exc

            if attempt < 3:
                prompt = (
                    base_prompt
                    + "\n\nRETRY REQUIREMENTS:\n"
                    + "- Use exact required headings only once each.\n"
                    + "- Remove placeholders and generic filler.\n"
                    + f"- Previous failure: {last_error}\n"
                )
        if last_error is None:
            raise RuntimeError("generation failed with unknown error")
        raise last_error

    def _build_prompt(self, rel_path: str, source_text: str, metadata: dict[str, Any]) -> str:
        system = (
            "You are Eidosian Scribe v2. Produce precise, source-grounded technical documentation. "
            "No placeholders, no guessing, no generic filler. If unknown, explicitly say insufficient evidence."
        )
        truncated = source_text[: self.cfg.max_chars]
        return (
            f"<|im_start|>system\n{system}<|im_end|>\n"
            f"<|im_start|>user\n"
            "Return Markdown with EXACT section headings:\n"
            "# File Overview\n"
            "## Key Structures\n"
            "## Behavior Summary\n"
            "## Validation Notes\n"
            "## Improvement Opportunities\n\n"
            f"Include source_path: `{rel_path}` and document_type: `{metadata.get('doc_type','unknown')}` in File Overview.\n"
            "In Validation Notes include at least 3 concrete verifiable facts quoted/paraphrased from source.\n"
            "In Key Structures include functions/classes/keys/events found in source.\n"
            "Keep content factual and concise.\n\n"
            f"SOURCE ({rel_path}):\n```\n{truncated}\n```\n"
            "<|im_end|>\n<|im_start|>assistant\n"
        )

    def _normalize(self, markdown: str) -> str:
        text = markdown.replace("\r\n", "\n").strip()
        if not text.startswith("# "):
            text = "# File Overview\n\n" + text
        missing = [heading for heading in REQUIRED_HEADINGS if heading not in text]
        if missing:
            raise ValueError(f"missing required headings: {missing}")
        return text + "\n"


class FederatedJudge:
    def __init__(self, cfg: ProcessorConfig) -> None:
        self.cfg = cfg

    def evaluate(self, *, markdown: str, source_text: str, rel_path: str, metadata: dict[str, Any]) -> dict[str, Any]:
        judges: list[dict[str, Any]] = []
        judges.append(self._judge_structure(markdown))
        judges.append(self._judge_safety(markdown))
        judges.append(self._judge_grounding(markdown, source_text))
        judges.append(self._judge_coverage(markdown, source_text))
        judges.append(self._judge_specificity(markdown))

        aggregate = round(sum(j["score"] for j in judges) / max(1, len(judges)), 4)
        min_score = min(j["score"] for j in judges) if judges else 0.0
        approved = aggregate >= self.cfg.approval_threshold and min_score >= 0.55

        return {
            "contract": "doc_forge.consensus_gate.v2",
            "evaluated_at": _now_iso(),
            "source_path": rel_path,
            "document_type": metadata.get("doc_type", "unknown"),
            "approved": approved,
            "aggregate_score": aggregate,
            "min_judge_score": round(min_score, 4),
            "judges": judges,
            "approval_threshold": self.cfg.approval_threshold,
        }

    def _judge_structure(self, markdown: str) -> dict[str, Any]:
        hits = sum(1 for h in REQUIRED_HEADINGS if h in markdown)
        score = hits / len(REQUIRED_HEADINGS)
        return {"name": "structure_contract", "score": round(score, 4), "details": {"required_hits": hits, "required_total": len(REQUIRED_HEADINGS)}}

    def _judge_safety(self, markdown: str) -> dict[str, Any]:
        lower = markdown.lower()
        hits = sorted({marker for marker in PLACEHOLDER_MARKERS if marker in lower})
        score = 1.0 if not hits else max(0.0, 1.0 - (0.18 * len(hits)))
        return {"name": "anti_placeholder", "score": round(score, 4), "details": {"markers": hits}}

    def _judge_grounding(self, markdown: str, source_text: str) -> dict[str, Any]:
        src_terms = set(_extract_terms(source_text, limit=80))
        doc_terms = set(_extract_terms(markdown, limit=200))
        overlap = len(src_terms & doc_terms)
        ratio = overlap / max(1, len(src_terms))
        score = min(1.0, ratio / max(0.01, self.cfg.min_grounding_overlap))
        return {
            "name": "grounding_overlap",
            "score": round(score, 4),
            "details": {"source_terms": len(src_terms), "overlap_terms": overlap, "ratio": round(ratio, 4)},
        }

    def _judge_coverage(self, markdown: str, source_text: str) -> dict[str, Any]:
        symbols = _extract_symbols(source_text)
        if not symbols:
            return {"name": "symbol_coverage", "score": 1.0, "details": {"symbols": 0, "hits": 0}}
        lower = markdown.lower()
        hits = sum(1 for s in symbols if s.lower() in lower)
        score = hits / max(1, min(20, len(symbols)))
        return {"name": "symbol_coverage", "score": round(min(1.0, score), 4), "details": {"symbols": len(symbols), "hits": hits}}

    def _judge_specificity(self, markdown: str) -> dict[str, Any]:
        lines = [line.strip() for line in markdown.splitlines() if line.strip()]
        bullet_lines = sum(1 for line in lines if line.startswith("- ") or line.startswith("* "))
        numbers = len(re.findall(r"\b\d+\b", markdown))
        evidence_markers = len(re.findall(r"`[^`]+`", markdown))
        raw = 0.4
        raw += min(0.25, bullet_lines * 0.01)
        raw += min(0.2, evidence_markers * 0.01)
        raw += min(0.15, numbers * 0.01)
        return {
            "name": "specificity_density",
            "score": round(min(1.0, raw), 4),
            "details": {"bullets": bullet_lines, "code_refs": evidence_markers, "numbers": numbers},
        }


class DocProcessor:
    def __init__(self, cfg: ProcessorConfig) -> None:
        self.cfg = cfg
        self.extractor = DocumentExtractor(cfg)
        self.generator = DocGenerator(cfg)
        self.judges = FederatedJudge(cfg)
        self.model_server = ManagedModelServer(cfg)

        self.stop_event = threading.Event()
        self.lock = threading.Lock()
        self.thread: threading.Thread | None = None
        self.processing_started = False

        self.excluded_prefixes = [
            (cfg.forge_root / ".git").resolve(),
            (cfg.forge_root / "eidosian_venv").resolve(),
            (cfg.forge_root / "llama.cpp").resolve(),
            (cfg.forge_root / "models").resolve(),
            (cfg.forge_root / "doc_forge" / "staging").resolve(),
            (cfg.forge_root / "doc_forge" / "final_docs").resolve(),
            cfg.runtime_root.resolve(),
        ]
        self.excluded_segments = {
            ".git", "__pycache__", ".pytest_cache", "node_modules", "archive_forge", "Backups", "runtime", "staging", "final_docs",
        }

        self.state = self._load_state()
        self.recent_docs: deque[dict[str, Any]] = deque(maxlen=80)

        for path in [
            cfg.runtime_root,
            cfg.staging_root,
            cfg.final_root,
            cfg.rejected_root,
            cfg.judgments_root,
        ]:
            path.mkdir(parents=True, exist_ok=True)

    def _default_state(self) -> dict[str, Any]:
        return {
            "started_at": _now_iso(),
            "model": str(self.cfg.llm_model_path),
            "status": "initializing",
            "total_discovered": 0,
            "processed": 0,
            "approved": 0,
            "rejected": 0,
            "errors": 0,
            "average_seconds_per_document": 0.0,
            "average_quality_score": 0.0,
            "last_staged": None,
            "last_approved": None,
            "tag_frequency": {},
            "theme_frequency": {},
            "doc_type_frequency": {},
            "files": {},
            "index": [],
            "history": [],
        }

    def _load_state(self) -> dict[str, Any]:
        state = _read_json(self.cfg.state_path, self._default_state())
        default = self._default_state()
        for key, value in default.items():
            state.setdefault(key, value)
        return state

    def _persist_state(self) -> None:
        with self.lock:
            snapshot = json.loads(json.dumps(self.state))
        _atomic_write_json(self.cfg.state_path, snapshot)
        _atomic_write_json(self.cfg.status_path, self.status_payload())

    def start(self) -> None:
        if self.thread and self.thread.is_alive():
            return
        self.stop_event.clear()
        self.thread = threading.Thread(target=self._run_loop, name="doc-processor", daemon=True)
        self.thread.start()

    def stop(self) -> None:
        self.stop_event.set()
        if self.thread:
            self.thread.join(timeout=10)
        self.model_server.stop()

    def _is_excluded(self, path: Path) -> bool:
        resolved = path.resolve()
        for prefix in self.excluded_prefixes:
            try:
                resolved.relative_to(prefix)
                return True
            except Exception:
                pass
        for part in resolved.parts:
            if part in self.excluded_segments:
                return True
        return False

    def _should_include(self, path: Path) -> bool:
        if not path.is_file():
            return False
        if self._is_excluded(path):
            return False
        suffix = path.suffix.lower()
        if suffix not in self.cfg.include_suffixes and suffix != "":
            return False
        try:
            size = path.stat().st_size
        except OSError:
            return False
        if size <= 0:
            return False
        if size > self.cfg.max_file_bytes and suffix not in {".pdf", ".docx"}:
            return False
        return True

    def _scan_candidates(self) -> list[Path]:
        root = self.cfg.source_root
        out: list[Path] = []
        for cur, dirs, files in os.walk(root):
            cur_path = Path(cur)
            dirs[:] = [d for d in dirs if not self._is_excluded(cur_path / d)]
            for filename in files:
                path = cur_path / filename
                if self._should_include(path):
                    out.append(path)
        out.sort()
        return out

    def _document_paths(self, rel_path: Path) -> tuple[Path, Path, Path]:
        stage = self.cfg.staging_root / rel_path.with_suffix(rel_path.suffix + ".md")
        final = self.cfg.final_root / rel_path.with_suffix(rel_path.suffix + ".md")
        rejected = self.cfg.rejected_root / rel_path.with_suffix(rel_path.suffix + ".md")
        return stage, final, rejected

    def _record_history(self, entry: dict[str, Any]) -> None:
        history = self.state.get("history", [])
        history.append(entry)
        self.state["history"] = history[-500:]
        self.recent_docs.appendleft(entry)

    def _update_running_average(self, key: str, value: float, count_key: str) -> None:
        current_count = int(self.state.get(count_key, 0))
        current_avg = float(self.state.get(key, 0.0))
        new_count = current_count + 1
        new_avg = ((current_avg * current_count) + value) / new_count
        self.state[key] = round(new_avg, 4)
        self.state[count_key] = new_count

    def _run_loop(self) -> None:
        with self.lock:
            self.state["status"] = "starting"
        self._persist_state()

        while not self.stop_event.is_set():
            if not self.cfg.dry_run and not self.model_server.is_ready():
                try:
                    self.model_server.start()
                    with self.lock:
                        self.state["status"] = "running"
                        self.state["last_error"] = ""
                except Exception as exc:
                    with self.lock:
                        self.state["status"] = "waiting_for_llm"
                        self.state["last_error"] = f"{type(exc).__name__}: {exc}"
                        self.state["last_error_ts"] = _now_iso()
                    self._persist_state()
                    time.sleep(max(2.0, self.cfg.loop_sleep_s))
                    continue
                self._persist_state()
            elif self.cfg.dry_run:
                status_changed = False
                with self.lock:
                    if self.state.get("status") == "starting":
                        self.state["status"] = "running"
                        status_changed = True
                if status_changed:
                    self._persist_state()

            candidates = self._scan_candidates()
            with self.lock:
                self.state["total_discovered"] = len(candidates)

            pending = []
            for path in candidates:
                rel = path.resolve().relative_to(self.cfg.source_root.resolve())
                rel_key = str(rel)
                raw = path.read_bytes()
                digest = _sha256_bytes(raw)
                file_state = self.state["files"].get(rel_key)
                if file_state and file_state.get("sha256") == digest and file_state.get("status") == "approved":
                    continue
                pending.append((path, rel, digest))

            if not pending:
                with self.lock:
                    self.state["status"] = "idle"
                self._persist_state()
                time.sleep(max(2.0, self.cfg.loop_sleep_s))
                with self.lock:
                    if self.state.get("status") == "idle":
                        self.state["status"] = "running"
                continue

            for path, rel, digest in pending:
                if self.stop_event.is_set():
                    break
                started = time.perf_counter()
                rel_key = str(rel)
                stage_path, final_path, rejected_path = self._document_paths(rel)
                judgment_path = self.cfg.judgments_root / rel.with_suffix(rel.suffix + ".json")
                for p in (stage_path, final_path, rejected_path, judgment_path):
                    p.parent.mkdir(parents=True, exist_ok=True)

                doc_type = rel.suffix.lower().lstrip(".") or "text"
                with self.lock:
                    self.state["doc_type_frequency"][doc_type] = self.state["doc_type_frequency"].get(doc_type, 0) + 1

                try:
                    source_text, metadata = self.extractor.extract(path)
                    metadata["source_size_bytes"] = path.stat().st_size
                    metadata["doc_type"] = metadata.get("doc_type") or doc_type
                    markdown = self.generator.generate(rel_key, source_text, metadata) if not self.cfg.dry_run else "\n".join(REQUIRED_HEADINGS)

                    stage_path.write_text(markdown, encoding="utf-8")
                    with self.lock:
                        self.state["last_staged"] = rel_key

                    scorecard = self.judges.evaluate(markdown=markdown, source_text=source_text, rel_path=rel_key, metadata=metadata)
                    _atomic_write_json(judgment_path, scorecard)

                    quality = float(scorecard.get("aggregate_score", 0.0))
                    self._update_running_average("average_quality_score", quality, "quality_samples")

                    tags = _extract_terms(markdown, limit=8)
                    themes = _extract_terms(source_text, limit=8)
                    with self.lock:
                        for t in tags:
                            self.state["tag_frequency"][t] = self.state["tag_frequency"].get(t, 0) + 1
                        for t in themes:
                            self.state["theme_frequency"][t] = self.state["theme_frequency"].get(t, 0) + 1

                    approved = bool(scorecard.get("approved", False))
                    status = "approved" if approved else "rejected"
                    if approved:
                        final_path.write_text(markdown, encoding="utf-8")
                        with self.lock:
                            self.state["approved"] += 1
                            self.state["last_approved"] = rel_key
                        index_entry = {
                            "source": rel_key,
                            "document": str(final_path.relative_to(self.cfg.runtime_root)),
                            "score": quality,
                            "updated_at": _now_iso(),
                            "tags": tags,
                            "doc_type": metadata.get("doc_type", "unknown"),
                        }
                        with self.lock:
                            current_index = [e for e in self.state.get("index", []) if e.get("source") != rel_key]
                            current_index.append(index_entry)
                            current_index.sort(key=lambda x: x.get("source", ""))
                            self.state["index"] = current_index
                    else:
                        rejected_path.write_text(markdown, encoding="utf-8")
                        with self.lock:
                            self.state["rejected"] += 1

                    elapsed = round(time.perf_counter() - started, 4)
                    self._update_running_average("average_seconds_per_document", elapsed, "duration_samples")

                    with self.lock:
                        self.state["processed"] += 1
                        self.state["files"][rel_key] = {
                            "sha256": digest,
                            "status": status,
                            "score": quality,
                            "doc_type": metadata.get("doc_type", "unknown"),
                            "updated_at": _now_iso(),
                            "staged_path": str(stage_path.relative_to(self.cfg.runtime_root)),
                            "final_path": str(final_path.relative_to(self.cfg.runtime_root)) if approved else None,
                            "rejected_path": str(rejected_path.relative_to(self.cfg.runtime_root)) if not approved else None,
                            "judgment_path": str(judgment_path.relative_to(self.cfg.runtime_root)),
                            "duration_seconds": elapsed,
                            "tags": tags,
                        }
                        self._record_history(
                            {
                                "source": rel_key,
                                "status": status,
                                "score": quality,
                                "duration_seconds": elapsed,
                                "updated_at": _now_iso(),
                            }
                        )
                except Exception as exc:
                    elapsed = round(time.perf_counter() - started, 4)
                    with self.lock:
                        self.state["processed"] += 1
                        self.state["errors"] += 1
                        self.state["files"][rel_key] = {
                            "sha256": digest,
                            "status": "error",
                            "error": str(exc),
                            "updated_at": _now_iso(),
                            "duration_seconds": elapsed,
                        }
                        self._record_history(
                            {
                                "source": rel_key,
                                "status": "error",
                                "error": str(exc),
                                "duration_seconds": elapsed,
                                "updated_at": _now_iso(),
                            }
                        )

                _atomic_write_json(self.cfg.index_path, {"entries": self.state.get("index", [])})
                self._persist_state()
                time.sleep(self.cfg.loop_sleep_s)

        with self.lock:
            self.state["status"] = "stopped"
        self._persist_state()

    def status_payload(self) -> dict[str, Any]:
        with self.lock:
            data = json.loads(json.dumps(self.state))

        total = int(data.get("total_discovered", 0))
        processed = int(data.get("processed", 0))
        remaining = max(0, total - processed)
        avg_seconds = float(data.get("average_seconds_per_document", 0.0))
        eta = round(remaining * avg_seconds, 2) if avg_seconds > 0 else None

        def _top(counter_map: dict[str, Any], n: int = 12) -> list[tuple[str, int]]:
            items = [(k, int(v)) for k, v in counter_map.items()]
            items.sort(key=lambda kv: (-kv[1], kv[0]))
            return items[:n]

        return StatusPayload(
            status=str(data.get("status", "unknown")),
            started_at=str(data.get("started_at", _now_iso())),
            model=str(data.get("model", "")),
            total_discovered=total,
            processed=processed,
            approved=int(data.get("approved", 0)),
            rejected=int(data.get("rejected", 0)),
            errors=int(data.get("errors", 0)),
            remaining=remaining,
            average_quality_score=float(data.get("average_quality_score", 0.0)),
            average_seconds_per_document=avg_seconds,
            eta_seconds=eta,
            last_staged=data.get("last_staged"),
            last_approved=data.get("last_approved"),
            top_tags=_top(data.get("tag_frequency", {})),
            top_themes=_top(data.get("theme_frequency", {})),
            doc_type_frequency=_top(data.get("doc_type_frequency", {})),
        ).model_dump()

    def recent_entries(self, limit: int = 50) -> list[dict[str, Any]]:
        with self.lock:
            items = list(self.state.get("history", []))
        return list(reversed(items[-limit:]))

    def index_entries(self, limit: int = 300) -> list[dict[str, Any]]:
        with self.lock:
            entries = list(self.state.get("index", []))
        return entries[:limit]


DASHBOARD_HTML = """<!doctype html>
<html lang=\"en\">
<head>
  <meta charset=\"utf-8\" />
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1\" />
  <title>Eidosian Scribe v2</title>
  <style>
    :root {
      --bg-0: #0b1116;
      --bg-1: #101a22;
      --panel: #162733;
      --line: #2f4759;
      --text: #e6f4ff;
      --sub: #90abc0;
      --ok: #00d084;
      --warn: #ffc857;
      --bad: #ff5d73;
      --accent: #5cc8ff;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      font-family: "Space Grotesk", "Segoe UI", sans-serif;
      color: var(--text);
      background:
        radial-gradient(1200px 500px at 10% -10%, #1d3140 0%, transparent 60%),
        radial-gradient(900px 420px at 100% 0%, #173145 0%, transparent 65%),
        linear-gradient(160deg, var(--bg-0), var(--bg-1));
      min-height: 100vh;
    }
    .topbar {
      position: sticky;
      top: 0;
      z-index: 10;
      background: rgba(14, 26, 34, 0.94);
      border-bottom: 1px solid var(--line);
      backdrop-filter: blur(6px);
      padding: 10px 16px;
      display: grid;
      grid-template-columns: repeat(8, minmax(120px, 1fr));
      gap: 8px;
    }
    .metric {
      background: rgba(26, 46, 60, 0.65);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 8px;
    }
    .metric .k { font-size: 11px; color: var(--sub); text-transform: uppercase; letter-spacing: 0.08em; }
    .metric .v { font-family: "JetBrains Mono", monospace; font-size: 14px; margin-top: 4px; }
    .wrap { max-width: 1400px; margin: 18px auto; padding: 0 16px 24px; }
    .grid { display: grid; grid-template-columns: 2fr 1fr; gap: 14px; }
    .card {
      background: rgba(24, 42, 56, 0.72);
      border: 1px solid var(--line);
      border-radius: 12px;
      overflow: hidden;
    }
    .card h2 { margin: 0; padding: 12px 14px; font-size: 14px; border-bottom: 1px solid var(--line); }
    table { width: 100%; border-collapse: collapse; font-size: 12px; }
    th, td { padding: 8px 10px; border-bottom: 1px solid rgba(47, 71, 89, 0.6); text-align: left; vertical-align: top; }
    th { color: var(--sub); font-weight: 500; }
    .pill { border-radius: 999px; padding: 2px 8px; font-size: 11px; display: inline-block; }
    .ok { background: rgba(0, 208, 132, 0.18); color: var(--ok); }
    .rej { background: rgba(255, 93, 115, 0.18); color: var(--bad); }
    .err { background: rgba(255, 200, 87, 0.2); color: var(--warn); }
    .mono { font-family: "JetBrains Mono", monospace; }
    .tags { display: flex; gap: 6px; flex-wrap: wrap; padding: 10px; }
    .tag { background: #243a49; border: 1px solid #355669; border-radius: 8px; padding: 4px 8px; font-size: 11px; }
    .small { font-size: 11px; color: var(--sub); }
  </style>
</head>
<body>
  <div class=\"topbar\" id=\"topbar\"></div>
  <div class=\"wrap\">
    <div class=\"grid\">
      <div class=\"card\">
        <h2>Documentation Index</h2>
        <table>
          <thead>
            <tr>
              <th>Source</th><th>Score</th><th>Type</th><th>Tags</th><th>Updated</th>
            </tr>
          </thead>
          <tbody id=\"indexRows\"></tbody>
        </table>
      </div>
      <div style=\"display:grid;gap:14px;\">
        <div class=\"card\">
          <h2>Top Tags</h2>
          <div id=\"tags\" class=\"tags\"></div>
        </div>
        <div class=\"card\">
          <h2>Top Themes</h2>
          <div id=\"themes\" class=\"tags\"></div>
        </div>
        <div class=\"card\">
          <h2>Recent Decisions</h2>
          <table>
            <thead><tr><th>Source</th><th>Status</th><th>Score</th><th>t(s)</th></tr></thead>
            <tbody id=\"recentRows\"></tbody>
          </table>
        </div>
      </div>
    </div>
  </div>
<script>
  const topbar = document.getElementById('topbar');
  const indexRows = document.getElementById('indexRows');
  const recentRows = document.getElementById('recentRows');
  const tags = document.getElementById('tags');
  const themes = document.getElementById('themes');

  function metric(k, v) {
    return `<div class=\"metric\"><div class=\"k\">${k}</div><div class=\"v\">${v}</div></div>`;
  }

  function fmtSec(v) {
    if (v === null || v === undefined) return 'n/a';
    if (v < 60) return `${v.toFixed(1)}s`;
    return `${(v/60).toFixed(1)}m`;
  }

  function fmtStatus(s) {
    if (s === 'approved') return '<span class=\"pill ok\">approved</span>';
    if (s === 'rejected') return '<span class=\"pill rej\">rejected</span>';
    if (s === 'error') return '<span class=\"pill err\">error</span>';
    return `<span class=\"pill\">${s}</span>`;
  }

  async function refresh() {
    const [statusRes, indexRes, recentRes] = await Promise.all([
      fetch('/api/status'),
      fetch('/api/index?limit=200'),
      fetch('/api/recent?limit=40')
    ]);
    const status = await statusRes.json();
    const index = await indexRes.json();
    const recent = await recentRes.json();

    topbar.innerHTML = [
      metric('Model', `<span class=\"mono\">${status.model.split('/').slice(-1)[0]}</span>`),
      metric('State', status.status),
      metric('Processed', `${status.processed}/${status.total_discovered}`),
      metric('Approved', status.approved),
      metric('Rejected', status.rejected),
      metric('Avg Score', Number(status.average_quality_score || 0).toFixed(3)),
      metric('Avg Time/Doc', fmtSec(status.average_seconds_per_document || 0)),
      metric('ETA', fmtSec(status.eta_seconds)),
      metric('Last Staged', `<span class=\"small\">${status.last_staged || 'n/a'}</span>`),
      metric('Last Approved', `<span class=\"small\">${status.last_approved || 'n/a'}</span>`),
    ].join('');

    indexRows.innerHTML = index.entries.map((entry) => `
      <tr>
        <td class=\"mono\">${entry.source}</td>
        <td>${Number(entry.score || 0).toFixed(3)}</td>
        <td>${entry.doc_type || ''}</td>
        <td class=\"small\">${(entry.tags || []).join(', ')}</td>
        <td class=\"small\">${entry.updated_at || ''}</td>
      </tr>
    `).join('');

    recentRows.innerHTML = recent.entries.map((entry) => `
      <tr>
        <td class=\"mono\">${entry.source}</td>
        <td>${fmtStatus(entry.status)}</td>
        <td>${entry.score !== undefined ? Number(entry.score).toFixed(3) : 'n/a'}</td>
        <td>${entry.duration_seconds !== undefined ? Number(entry.duration_seconds).toFixed(2) : 'n/a'}</td>
      </tr>
    `).join('');

    tags.innerHTML = (status.top_tags || []).map(([k,v]) => `<div class=\"tag\">${k} <span class=\"small\">${v}</span></div>`).join('');
    themes.innerHTML = (status.top_themes || []).map(([k,v]) => `<div class=\"tag\">${k} <span class=\"small\">${v}</span></div>`).join('');
  }

  refresh();
  setInterval(refresh, 2500);
</script>
</body>
</html>
"""


def create_app(processor: DocProcessor) -> FastAPI:
    app = FastAPI(title="Eidosian Doc Processor", version="2.0.0")

    @app.on_event("startup")
    def _startup() -> None:
        processor.start()

    @app.on_event("shutdown")
    def _shutdown() -> None:
        processor.stop()

    @app.get("/health")
    def health() -> JSONResponse:
        return JSONResponse({"status": "ok", "service": "doc_processor", "time": _now_iso()})

    @app.get("/api/status")
    def api_status() -> JSONResponse:
        return JSONResponse(processor.status_payload())

    @app.get("/api/index")
    def api_index(limit: int = Query(default=300, ge=1, le=2000)) -> JSONResponse:
        return JSONResponse({"entries": processor.index_entries(limit=limit)})

    @app.get("/api/recent")
    def api_recent(limit: int = Query(default=60, ge=1, le=500)) -> JSONResponse:
        return JSONResponse({"entries": processor.recent_entries(limit=limit)})

    @app.get("/", response_class=HTMLResponse)
    def dashboard() -> str:
        return DASHBOARD_HTML

    return app


def main() -> int:
    parser = argparse.ArgumentParser(description="Run Eidosian Doc Processor v2 service")
    parser.add_argument("--host", default=os.environ.get("EIDOS_DOC_FORGE_HOST", "127.0.0.1"))
    parser.add_argument("--port", type=int, default=None)
    parser.add_argument("--dry-run", action="store_true", help="Skip model generation and emit heading shell only")
    parser.add_argument(
        "--forge-root",
        default=os.environ.get("EIDOS_FORGE_ROOT", "/data/data/com.termux/files/home/eidosian_forge"),
    )
    args = parser.parse_args()

    forge_root = Path(args.forge_root).resolve()
    default_doc_port = _registry_port(forge_root, "doc_forge_dashboard", 8930)
    effective_port = args.port if args.port is not None else int(os.environ.get("EIDOS_DOC_FORGE_PORT", str(default_doc_port)))
    cfg = ProcessorConfig.from_env(forge_root=forge_root, host=args.host, port=effective_port, dry_run=bool(args.dry_run))
    processor = DocProcessor(cfg)
    app = create_app(processor)

    uvicorn.run(app, host=cfg.host, port=cfg.port, log_level=os.environ.get("EIDOS_DOC_FORGE_LOG_LEVEL", "info"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
