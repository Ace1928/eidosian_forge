import os
import json
import csv
import xml.etree.ElementTree as ET
from typing import Any, Dict, Union, Optional, Tuple, List
import pandas as pd
import logging
import yaml
import pickle
import configparser
import markdown
import openpyxl
import sqlite3
import PyPDF2
import PIL.Image
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, Label, Toplevel
from PIL import Image, ImageTk
import os
import logging
import json
import pandas as pd
def create_widgets(self):
    top_frame = ttk.Frame(self.master)
    top_frame.pack(fill=tk.X, padx=10, pady=5)
    middle_frame = ttk.Frame(self.master)
    middle_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
    bottom_frame = ttk.Frame(self.master)
    bottom_frame.pack(fill=tk.X, padx=10, pady=5)
    path_label = ttk.Label(top_frame, text='File Path:')
    path_label.pack(side=tk.LEFT, padx=(0, 10))
    path_entry = ttk.Entry(top_frame, textvariable=self.file_path, width=60)
    path_entry.pack(side=tk.LEFT, expand=True, fill=tk.X)
    browse_button = ttk.Button(top_frame, text='Browse', command=self.browse_file)
    browse_button.pack(side=tk.RIGHT, padx=(10, 5))
    read_button = ttk.Button(top_frame, text='Read File', command=self.read_file)
    read_button.pack(side=tk.RIGHT)
    self.tab_parent = ttk.Notebook(middle_frame)
    self.tab_parent.pack(expand=1, fill='both')
    self.data_frame_tab = ttk.Frame(self.tab_parent)
    self.tab_parent.add(self.data_frame_tab, text='Data')
    self.image_frame_tab = ttk.Frame(self.tab_parent)
    self.tab_parent.add(self.image_frame_tab, text='Image')
    self.text_frame_tab = ttk.Frame(self.tab_parent)
    self.tab_parent.add(self.text_frame_tab, text='Text')
    self.data_tree = ttk.Treeview(self.data_frame_tab, columns=('Data',), show='headings')
    self.data_tree.pack(expand=True, fill='both')
    self.image_label = Label(self.image_frame_tab)
    self.image_label.pack(expand=True, fill='both')
    self.text_widget = tk.Text(self.text_frame_tab, wrap='word')
    self.text_widget.pack(expand=True, fill='both')
    self.status_label = ttk.Label(bottom_frame, text='Select a file to get started.', relief=tk.SUNKEN, anchor=tk.W)
    self.status_label.pack(fill=tk.X)